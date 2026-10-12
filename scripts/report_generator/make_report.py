#!/usr/bin/env python3
"""Empire FGA scouting report generator (workbook -> HTML -> PDF).

Usage:
  python3 make_report.py --workbook WB.xlsx --bio bio.json --pairing pairing.json \
      --out OUTDIR [--rule all|positive] [--final]

Scores are recomputed from the EVENT_LOG sheet. The cached SCORES sheet is only used
as a cross-check. Anything unverified makes the output a DRAFT; the QA file lists why.
"""
import argparse, collections, html, json, os, sys
import openpyxl

OBI = {"OBI_T": .35, "OBI_S": .35, "OBI_P": .30}
SII = {"SII_A": .40, "SII_T": .35, "SII_R": .25}
NAMES = {
 "OBI_T": "Run Timing", "OBI_S": "Space Architecture", "OBI_P": "Positional Anticipation",
 "SII_A": "Explosive Acceleration", "SII_T": "Top-End Speed Relative", "SII_R": "Recovery Sprint Frequency",
 "DMA": "Decision Making", "CHR": "Character", "MCS": "Mental Competence", "SOA": "Speed on Action",
 "SPI": "Spatial Performance", "CLM": "Calmness", "CRS": "Clutch Response", "OMS": "Transition",
 "DYN": "Dynamism", "CRK": "Creative Risk",
 "DRB": "Dribbling", "INT": "Interceptions", "KEY": "Key Passes", "TAC": "Tackling", "STK": "Shot Technique",
 "LGP": "Long Pass", "OTP": "One-Touch Pass", "DGP": "Diagonal Pass", "HZP": "Horizontal Pass",
 "UNP": "Unexpected Pass", "VRP": "Vertical Pass", "CAR": "Ball Carries", "RDS": "Space Reading", "MRK": "Marking",
}
REMAINING = ["DMA", "CHR", "MCS", "SOA", "SPI", "CLM", "CRS", "OMS", "DYN", "CRK"]
SUPPORT = ["DRB", "INT", "KEY", "TAC", "STK", "LGP", "OTP", "DGP", "HZP", "UNP", "VRP", "CAR"]
CORE_SUBS = list(OBI) + list(SII) + REMAINING


def esc(s):
    return html.escape(str(s), quote=True)


def load(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    prof = wb["PLAYER_PROFILE"]
    matches = []
    for r in prof.iter_rows(min_row=1, max_row=40, values_only=True):
        if r[1] is not None and str(r[1]).strip().isdigit() and r[3] is not None and r[4] is not None:
            matches.append(dict(n=int(r[1]), date=r[2], ha=str(r[3]), opp=str(r[4]),
                                state=str(r[5] or ""), tier=str(r[6] or "")))
    events = []
    for r in wb["EVENT_LOG"].iter_rows(min_row=4, values_only=True):
        if r[2] and r[5] is not None and r[1] is not None:
            try:
                events.append(dict(m=int(float(str(r[1]))), metric=str(r[2]).strip(), ts=str(r[3]),
                                   desc=str(r[4] or "").strip(), q=float(r[5]), typ=str(r[6] or "")))
            except ValueError:
                pass
    cached = {}
    for r in wb["SCORES"].iter_rows(min_row=5, max_row=70, values_only=True):
        if r[1] and isinstance(r[1], str) and " — " in r[1]:
            code = r[1].split(" — ")[0].strip()
            if code not in cached and isinstance(r[7], (int, float)):
                cached[code] = r[7]
    meta = dict(
        analyst=str(prof["C9"].value or ""), started=str(prof["C10"].value or "")[:10],
        reason=str(prof["C12"].value or ""))
    return matches, events, cached, meta


def compute(events, rule):
    by = collections.defaultdict(list)
    for e in events:
        by[e["metric"]].append(e)
    res = {}
    for k, v in by.items():
        use = v if rule == "all" else [e for e in v if e["typ"].startswith("Positive")]
        n = len(v)
        matches = len({e["m"] for e in v})
        score = sum(e["q"] for e in use) / len(use) if use else None
        ok = n >= 3 and matches >= 3
        res[k] = dict(n=n, matches=matches, score=score, ok=ok,
                      anti=sum(1 for e in v if not e["typ"].startswith("Positive")))
    return res


def composite(res, weights):
    if all(k in res and res[k]["ok"] and res[k]["score"] is not None for k in weights):
        return sum(res[k]["score"] * w for k, w in weights.items())
    return None


def band(x):
    if x is None:
        return "—"
    return "ELITE" if x >= 9 else "HIGH" if x >= 7 else "COMPETENT" if x >= 5 else "DEVELOPING"


def f2(x):
    return "—" if x is None else f"{x:.2f}"


def badge(kind, text):
    return f'<span class="b {kind}">{esc(text)}</span>'


def metric_row(code, r, bold=True):
    if r is None:
        return f'<tr><td>{code} — {NAMES.get(code, code)}</td><td>0</td><td>0</td><td>—</td><td>—</td><td>{badge("pend", "NOT LOGGED")}</td></tr>'
    conf = badge("ok", "OK") if r["ok"] else badge("pend", f"VIDEO_PENDING ({r['matches']} match{'es' if r['matches'] != 1 else ''})")
    sc = f2(r["score"]) if r["ok"] else f2(r["score"])
    cls = "" if r["ok"] else ' class="dim"'
    return (f'<tr{cls}><td>{code} — {NAMES.get(code, code)}</td><td>{r["n"]}</td><td>{r["matches"]}</td>'
            f'<td><b>{sc}</b></td><td>{band(r["score"])}</td><td>{conf}</td></tr>')


def table(rows, head=("SUB-METRIC", "EVENTS", "MATCHES", "SCORE", "BAND", "CONFIDENCE")):
    h = "".join(f"<th>{x}</th>" for x in head)
    return f'<table class="t"><tr>{h}</tr>{"".join(rows)}</table>'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--bio", required=True)
    ap.add_argument("--pairing", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--rule", choices=["all", "positive"], default="all")
    ap.add_argument("--final", action="store_true")
    ap.add_argument("--version", default="v2.3")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    matches, events, cached, meta = load(a.workbook)
    bio = json.load(open(a.bio, encoding="utf-8"))
    pairing = json.load(open(a.pairing, encoding="utf-8"))
    res = compute(events, a.rule)
    obi = composite(res, OBI)
    sii = composite(res, SII)

    qa = []  # items that block FINAL status
    info = []  # items for Jim only

    # ---- cross-checks -------------------------------------------------------
    alt = compute(events, "positive" if a.rule == "all" else "all")
    for k, r in res.items():
        if k in cached and r["score"] is not None:
            pos = compute(events, "positive")[k]["score"]
            if pos is not None and abs(cached[k] - pos) > 0.01:
                qa.append(f"SCORES tab for {k} shows {cached[k]:.3f} but EVENT_LOG gives {pos:.3f} (positive-only). Cached value is stale; the report uses the event log.")
    mis = [e for e in events if e["q"] <= 3 and e["typ"].startswith("Positive")]
    if mis:
        qa.append(f"{len(mis)} event(s) with quality <=3 are tagged 'Positive indicator' (" + "; ".join(f"M{e['m']} {e['metric']} {e['ts']} q{e['q']:g}" for e in mis) + "). Tagging decides which events the SCORES tab averages.")
    diffs = [(k, res[k]["score"], alt[k]["score"]) for k in CORE_SUBS if k in res and res[k]["score"] is not None and alt.get(k) and alt[k]["score"] is not None and abs(res[k]["score"] - alt[k]["score"]) > 0.005]
    if diffs:
        qa.append(f"Averaging rule matters: with rule '{a.rule}' vs the other rule, {len(diffs)} core metrics differ (" + ", ".join(f"{k} {x:.2f} vs {y:.2f}" for k, x, y in diffs) + "). Jim must rule on which applies. This draft uses '" + a.rule + "'.")
    per_match = collections.Counter(e["m"] for e in events)
    thin = [m for m in matches if per_match.get(m["n"], 0) < 10]
    for m in thin:
        qa.append(f"Match {m['n']} ({m['opp']}) has only {per_match.get(m['n'], 0)} logged events but counts toward the 5-match rule.")
    unv = [p["name"] for p in pairing["players"] if not p.get("verified")]
    if unv:
        qa.append("Comparison heights/feet not verified against Transfermarkt: " + ", ".join(unv))
    if not bio.get("photo"):
        info.append("No player photo supplied. v2.2 used a YouTube channel avatar; confirm usage rights before adding one.")
    if not bio.get("market_value"):
        info.append("Market value omitted: v2.2 used FMInside (a Football Manager database). Supply a Transfermarkt value to add the row.")
    years = sorted({str(m["date"])[:4] for m in matches})
    info.append("Match dates span " + ", ".join(years) + "; report shows the register so the reader can see the seasons.")
    info.append("Band cutoffs (ELITE >=9, HIGH >=7, COMPETENT >=5) are inferred from the v2.2 report; no written rubric defines them. Confirm or supply the official cutoffs.")
    info.append("Tier 1 and Tier 2 not populated; Tier 1 value 5.50 in v2.2 has no source in the workbook, so it is not repeated.")

    draft = bool(qa) or not a.final
    status_chip = "TIER 3 SUBSTANTIVE · FGA_V3_FINAL NOT ISSUED (PENDING TIER 2)"

    # ---- sections -----------------------------------------------------------
    completeness = 0
    pend_core = []
    if obi is not None:
        completeness += 1
    else:
        pend_core.append("OBI")
    if sii is not None:
        completeness += 1
    else:
        pend_core.append("SII")
    for k in REMAINING:
        if k in res and res[k]["ok"]:
            completeness += 1
        else:
            pend_core.append(k)
    graded = [obi, sii] + [res[k]["score"] for k in REMAINING if k in res and res[k]["ok"]]
    graded = [x for x in graded if x is not None]
    t3avg = sum(graded) / len(graded) if graded else None

    opp = {m["n"]: m["opp"] for m in matches}
    top3 = sum(1 for m in matches if m["tier"].startswith("1"))
    away = sum(1 for m in matches if m["ha"].lower() == "away")
    losing = sum(1 for m in matches if m["state"].lower() == "losing")
    ok_sub = sum(1 for k in CORE_SUBS if k in res and res[k]["ok"])

    scorecard = f"""
<div class="cards">
 <div class="card"><div class="cl">VIDEO COMPLETENESS</div><div class="cv g">{completeness} / 12</div><div class="cs">{'Only ' + ', '.join(pend_core) + ' pending' if pend_core else 'All core metrics graded'}</div></div>
 <div class="card"><div class="cl">OBI COMPOSITE</div><div class="cv g">{f2(obi)}</div><div class="cs">{'Transfer ready' if obi is not None else 'Pending'}</div></div>
 <div class="card"><div class="cl">SII COMPOSITE</div><div class="cv a">{f2(sii)}</div><div class="cs">{'Transfer ready' if sii is not None else 'Pending'}</div></div>
</div>"""

    obi_rows = [metric_row(k, res.get(k)) for k in OBI] + [
        f'<tr class="comp"><td>OBI COMPOSITE</td><td>—</td><td>—</td><td>{f2(obi)}</td><td>{band(obi)}</td><td>{badge("ok", "TRANSFER READY") if obi is not None else badge("pend", "PENDING")}</td></tr>']
    sii_rows = [metric_row(k, res.get(k)) for k in SII] + [
        f'<tr class="comp"><td>SII COMPOSITE</td><td>—</td><td>—</td><td>{f2(sii)}</td><td>{band(sii)}</td><td>{badge("ok", "TRANSFER READY") if sii is not None else badge("pend", "PENDING")}</td></tr>']
    rem_rows = [metric_row(k, res.get(k)) for k in REMAINING]
    sup_rows = [metric_row(k, res.get(k)) for k in SUPPORT if k in res]

    # pairing
    P = pairing["players"]
    subj_h, subj_f = pairing["subject_height_cm"], pairing["subject_foot"]
    order = {"primary": 0, "secondary": 1, "tertiary": 2, "reference": 3, "demoted": 4, "excluded": 5}
    P.sort(key=lambda p: order[p["role"]])
    prow = []
    for p in P:
        gap = abs(p["height_cm"] - subj_h)
        passes = p["foot"] == subj_f and gap <= 5
        prow.append(f'<tr><td><b>{esc(p["name"])}</b> <span class="dim">({esc(p["club"])})</span></td><td>{p["height_cm"]} cm</td><td>{esc(p["foot"])}</td><td>{gap} cm</td>'
                    f'<td>{badge("ok", "PASSES FILTER") if passes else badge("fail", "FAILS FILTER")}</td><td>{esc(p["role"].upper())}</td></tr>')
    notes = "".join(f'<p><b>{esc(p["role"].upper())} — {esc(p["name"])}.</b> {esc(p["note"])}</p>' for p in P if p["role"] in ("primary", "secondary", "tertiary", "reference"))
    notes_dem = "".join(f'<p><b>{esc(p["role"].upper())} — {esc(p["name"])}.</b> {esc(p["note"])}</p>' for p in P if p["role"] in ("demoted", "excluded"))
    tree = "".join(
        f'<div class="node {"n1" if p["role"] in ("primary", "secondary", "tertiary") else "n2"}"><b>{esc(p["name"].upper())}</b><br><span>{esc(p["club"])}<br>{p["height_cm"]}cm · {esc(p["foot"])}-footed · gap {abs(p["height_cm"] - subj_h)}cm<br><i>{esc(p["role"]).upper()}</i></span></div>'
        for p in P if p["role"] != "excluded")

    # defining incidents
    groups = collections.defaultdict(list)
    for e in events:
        groups[(e["m"], e["ts"])].append(e)
    scored = []
    for (m, ts), v in groups.items():
        desc = max((e["desc"] for e in v), key=len)
        if len(desc) < 25:
            continue
        scored.append((sum(e["q"] for e in v) / len(v), len(v), m, ts, v, desc))
    best = sorted([s for s in scored if s[1] >= 2 and s[0] >= 9.0], key=lambda s: (-s[0], -s[1]))[:5]
    worst = sorted(scored, key=lambda s: (s[0], -s[1]))[:1]
    inc = ""
    for kind, lst in (("g", best), ("r", worst)):
        for avg, n, m, ts, v, desc in lst:
            sc = ", ".join(f"{e['q']:g}/10 on {e['metric']}" for e in sorted(v, key=lambda e: -e["q"]))
            inc += f'<div class="q {kind}"><b>M{m} {esc(ts)} (vs {esc(opp.get(m, "?"))}):</b> <i>“{esc(desc)}”</i> — scored <b>{esc(sc)}</b>.</div>'
    lowest = min((k for k in REMAINING if k in res and res[k]["ok"]), key=lambda k: res[k]["score"], default=None)
    lowest_block = ""
    if lowest:
        evs = sorted([e for e in events if e["metric"] == lowest], key=lambda e: (e["m"], e["ts"]))
        listing = ", ".join(f"{int(e['q']) if e['q'] == int(e['q']) else e['q']}/10 at {esc(e['ts'])} M{e['m']}" for e in evs)
        quote = f' The analyst noted: <i>“{esc(bio["analyst_note"])}”</i>' if bio.get("analyst_note") and lowest == bio.get("analyst_note_metric") else ""
        lowest_block = f'<div class="q bl"><b>The lowest-scoring core metric: {lowest} — {NAMES[lowest]} ({res[lowest]["score"]:.2f} across {res[lowest]["matches"]} matches).</b> Events: {listing}.{quote}</div>'

    # risk vs reward
    pos, cau = [], []
    for k in CORE_SUBS + ["DRB", "INT"]:
        r = res.get(k)
        if r and r["ok"] and r["score"] is not None:
            if r["score"] >= 8.0:
                pos.append(f"<b>{k} {r['score']:.2f} across {r['matches']} matches</b> — {NAMES[k].lower()}, {band(r['score']).lower()}")
            elif r["score"] < 6.0:
                cau.append(f"<b>{k} {r['score']:.2f} across {r['matches']} matches</b> — {NAMES[k].lower()} below competent range")
            if r["anti"] and r["score"] >= 6.0:
                cau.append(f"<b>{k}</b> has {r['anti']} anti-indicator event(s) in the log")
    for k in REMAINING:
        if k in res and not res[k]["ok"]:
            cau.append(f"<b>{k} VIDEO_PENDING</b> — {res[k]['n']} event(s) in {res[k]['matches']} match(es)")
    pend_sup = [k for k in SUPPORT if k in res and not res[k]["ok"]]
    if pend_sup:
        cau.append("<b>Supporting metrics below 3 matches:</b> " + ", ".join(pend_sup))
    if sii is not None and sii < 7:
        cau.append(f"<b>SII composite {sii:.2f}</b> — competent sprint profile, not elite")
    cau += [esc(x) for x in bio.get("caution_extra", [])]
    pos += [esc(x) for x in bio.get("positive_extra", [])]
    li = lambda L: "".join(f"<li>{x}</li>" for x in L)

    # validation
    v = []
    def vrow(req, ok, text):
        v.append(f'<tr><td>{req}</td><td>{badge("ok" if ok is True else "fail" if ok is False else "pend", text)}</td></tr>')
    vrow("Minimum 5 matches registered", len(matches) >= 5, f"{len(matches)} MATCHES")
    vrow("≥2 vs top-third opposition", top3 >= 2, str(top3))
    vrow("≥1 away match", away >= 1, str(away))
    vrow("≥1 match while losing", losing >= 1, str(losing))
    vrow("3 events in 3 distinct matches per core sub-metric", ok_sub == len(CORE_SUBS), f"{ok_sub} OF {len(CORE_SUBS)} PASS")
    vrow("Video completeness (of 12)", completeness == 12, f"{completeness}/12")
    vrow("Matches with at least 10 logged events", all(per_match.get(m["n"], 0) >= 10 for m in matches), f"{sum(1 for m in matches if per_match.get(m['n'], 0) >= 10)} OF {len(matches)}")
    vrow("COACH metric", bool(bio.get("coach")), "SCORED" if bio.get("coach") else "NOT SCORED")
    vrow("Second scout", bool(bio.get("second_scout")), "ASSIGNED" if bio.get("second_scout") else "NOT ASSIGNED")
    vrow("Tier 2 composite", False, "NOT POPULATED")
    vrow("Tier 1 statistical foundation", False, "NOT POPULATED")
    vrow("FGA_v3_Final", None, "NOT ISSUED — PENDING TIER 2")
    mrows = "".join(
        f'<tr><td>{m["n"]}</td><td>{esc(str(m["date"])[:10])}</td><td>{esc(m["ha"])}</td><td>{esc(m["opp"])}</td><td>{esc(m["state"])}</td><td>{esc(m["tier"])}</td><td>{per_match.get(m["n"], 0)}</td></tr>' for m in matches)

    rule_txt = ("Each sub-metric score is the mean quality (1–10) of all logged events for that sub-metric, anti-indicator events included."
                if a.rule == "all" else
                "Each sub-metric score is the mean quality (1–10) of events tagged 'Positive indicator'; anti-indicator events are excluded.")

    prof_rows = [("Full name", bio["full_name"]), ("Position", bio["position"]), ("Club", bio["club"]), ("League", bio["league"]),
                 ("Season", bio["season"]), ("Age", bio["age_line"]), ("Height", bio["height"]), ("Foot", bio["foot"]),
                 ("Nationality", bio["nationality"]), ("Contract", bio["contract"])]
    if bio.get("market_value"):
        prof_rows.append(("Market value", bio["market_value"]))
    prof_rows += [("Youth career", bio["youth"]), ("Senior career", bio["senior"]),
                  ("Tier 3 evidence", f"{len(matches)} matches · {len(events)} logged events"),
                  ("Video completeness", f"{completeness}/12"), ("COACH metric", "Scored" if bio.get("coach") else "Not scored"),
                  ("Second scout", "Assigned" if bio.get("second_scout") else "Not assigned")]
    prof = "".join(f"<tr><td>{esc(k)}</td><td>{esc(x)}</td></tr>" for k, x in prof_rows)

    t3line = f"{t3avg:.2f}" if t3avg is not None else "—"
    sources = "".join(f"<tr><td>{esc(s[0])}</td><td>{esc(s[1])}</td></tr>" for s in bio["sources"])
    wm = '<div class="wm">DRAFT</div>' if draft else ""
    ref = f'{bio["ref"]} ({a.version})'

    css = """
@page{size:A4;margin:16mm 15mm}
*{box-sizing:border-box}
body{font-family:Helvetica,Arial,sans-serif;color:#1a1a1a;font-size:10.4pt;line-height:1.45}
h1{font-size:25pt;margin:0 0 6px}
h2{font-size:12.5pt;color:#1F3A8A;text-transform:uppercase;border-bottom:1px solid #ddd;padding-bottom:5px;margin:24px 0 10px;page-break-after:avoid}
h3{font-size:10.5pt;text-transform:uppercase;margin:14px 0 6px;page-break-after:avoid}
.top{display:flex;justify-content:space-between;font-size:8pt;letter-spacing:.14em;color:#555;text-transform:uppercase}
.chip{border:2px solid #1b7a4b;color:#1b7a4b;font-weight:700;letter-spacing:.1em;font-size:8.5pt;padding:8px 12px;display:inline-block;margin:10px 0 4px}
.rule{border-top:3px solid #1F3A8A;margin:10px 0 12px}
table{border-collapse:collapse;width:100%}
.t th,.kv th{font-size:8pt;text-transform:uppercase;letter-spacing:.06em;color:#666;text-align:left;border-bottom:2px solid #1F3A8A;padding:6px 6px}
.t td,.kv td{border-bottom:1px solid #eee;padding:6px 6px;font-size:9.6pt;vertical-align:top}
.kv td:first-child{width:30%}
tr.comp td{font-weight:700}tr.dim td,.dim{color:#777}
.b{display:inline-block;font-size:8pt;font-weight:700;letter-spacing:.04em;padding:2px 7px;border-radius:3px;border:1px solid}
.b.ok{color:#1b7a4b;border-color:#9fd1b6;background:#f2faf5}
.b.pend{color:#9a6a00;border-color:#e6cf8f;background:#fffaf0}
.b.fail{color:#b3261e;border-color:#e9b3af;background:#fff4f3}
.cards{display:flex;gap:10px;margin:8px 0 14px}
.card{flex:1;border:1px solid #bfe3d4;border-radius:6px;padding:12px;text-align:center}
.cl{font-size:8pt;letter-spacing:.1em;color:#555}.cv{font-size:25pt;font-weight:700}.cv.g{color:#1b7a4b}.cv.a{color:#8a6d00}.cs{font-size:8.5pt;color:#555}
.q{border-left:4px solid #1b7a4b;padding:8px 12px;margin:9px 0;background:#fafafa;page-break-inside:avoid}
.q.r{border-color:#b3261e}.q.bl{border-color:#1F3A8A}
.two{display:flex;gap:14px}.two>div{flex:1}
ul{margin:4px 0 4px 16px;padding:0}li{margin:3px 0;font-size:9.6pt}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:8px}
.node{border:1px solid #d9b44a;background:#fffaf0;padding:8px 10px;font-size:8.6pt;text-align:center;border-radius:4px}
.node.n2{border-color:#e9b3af;background:#fff4f3}
.note{font-size:8.6pt;color:#555;margin-top:6px}
.wm{position:fixed;top:42%;left:12%;font-size:130pt;color:rgba(180,0,0,.07);transform:rotate(-30deg);z-index:-1;font-weight:700}
.foot{margin-top:22px;font-size:8.4pt;color:#555;border-top:1px solid #ddd;padding-top:8px}
.pb{page-break-before:always}
"""
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>Empire FGA Scouting Report — {esc(bio['name'])} {a.version}</title><style>{css}</style></head><body>{wm}
<div class="top"><span><b>ESN · Empire Scouting Network</b></span><span style="text-align:right">SCOUTING REPORT · {a.version.upper()}{' · DRAFT' if draft else ''}<br>REF: {esc(ref)}<br>{esc(bio['header_line'])}</span></div>
<h1>{esc(bio['name'].upper())}</h1>
<div>{esc(bio['subhead1'])}<br>{esc(bio['subhead2'])}<br><b>Tier 3: {len(events)} events · {len(matches)} matches</b> · Video completeness <b>{completeness}/12</b><br>{('<a href="' + esc(bio['video_url']) + '">Watch video highlights on YouTube</a>') if bio.get('video_url') else ''}</div>
<div class="chip">{status_chip}</div><div class="rule"></div>

<h2>1 · Subject profile</h2><table class="kv">{prof}</table>
<p class="note">{esc(bio['log_note'])}</p>
<h3>Match register</h3>
<table class="t"><tr><th>#</th><th>Date</th><th>H/A</th><th>Opponent</th><th>Team state</th><th>Opposition</th><th>Events logged</th></tr>{mrows}</table>

<h2>2 · Tier 3 event log — scorecard</h2>{scorecard}
<h3>OBI — Off-Ball Intelligence</h3>{table(obi_rows)}
<h3>SII — Sprint Intensity Index</h3>{table(sii_rows)}
<h3>Remaining Tier 3 metrics</h3>{table(rem_rows)}
<h3>Supporting metrics</h3>{table(sup_rows)}
<p class="note"><b>Method.</b> {esc(rule_txt)} A sub-metric is confirmed only with 3+ events in 3+ distinct matches; otherwise it is shown as VIDEO_PENDING and excluded from composites.</p>

<h2 class="pb">3 · Tactical DNA pairing — foot and height filter applied</h2>
<p>{esc(pairing['filter_text'])}</p>
<table class="t"><tr><th>Reference player</th><th>Height</th><th>Foot</th><th>Gap vs subject ({subj_h} cm)</th><th>Filter</th><th>Role</th></tr>{''.join(prow)}</table>
{notes}{notes_dem}

<h2>4 · Genealogical tree — filtered lineage</h2>
<div class="grid">{tree}</div>
<p class="note">{esc(pairing['tree_note'])}</p>

<h2 class="pb">5 · The defining incidents — analyst quotes</h2>{inc}{lowest_block}

<h2>6 · FGA rating — status</h2>
<table class="t"><tr><th>Tier</th><th>Weight</th><th>Status</th><th>Value</th></tr>
<tr><td>Tier 3 — Scout Intelligence</td><td>50%</td><td>{badge('ok', f'{completeness}/12 CORE METRICS GRADED')}</td><td>{t3line} (indicative core average)</td></tr>
<tr><td>Tier 2 — Composite</td><td>30%</td><td>{badge('fail', 'NOT POPULATED')}</td><td>—</td></tr>
<tr><td>Tier 1 — Statistical Foundation</td><td>20%</td><td>{badge('fail', 'NOT POPULATED')}</td><td>—</td></tr>
<tr><td><b>FGA_v3_Final</b></td><td>—</td><td>{badge('pend', 'NOT ISSUED — PENDING TIER 2')}</td><td>—</td></tr></table>
<p class="note">The Tier 3 core average is the mean of the {len(graded)} confirmed core values (OBI, SII and the named sub-metrics at full confidence). It is an indicator, not an FGA rating.</p>

<h2>7 · Risk vs reward</h2>
<div class="two"><div><h3>Positive filters</h3><ul>{li(pos)}</ul></div><div><h3>Caution signals &amp; gaps</h3><ul>{li(cau)}</ul></div></div>

<h2 class="pb">8 · Validation status</h2><table class="kv">{''.join(v)}</table>

<h2>9 · Market reality</h2><table class="kv">{''.join(f'<tr><td>{esc(k)}</td><td>{esc(x)}</td></tr>' for k, x in bio['market_rows'])}</table>
<p>{esc(bio['market_text'])}</p>
<h2>Appendix — source inventory</h2><table class="kv">{sources}</table>
<div class="foot">Brandsen Sports · The Hague, Netherlands · jim.totime@brandsensports.com · empirescoutingnetwork.com<br>Report ref {esc(ref)} · Method: Empire FGA OBI &amp; SII Rubric Framework v1.0 · 50/30/20 master weighting · LOG is a context badge, not a multiplier<br><i>Confidential. Proprietary to Empire Scouting Network. Generated from workbook data on {esc(bio['generated'])}.</i></div>
</body></html>"""

    stem = f"Empire_FGA_Scouting_Report_{bio['name'].replace(' ', '_')}_{a.version}{'_DRAFT' if draft else ''}"
    hp = os.path.join(a.out, stem + ".html")
    open(hp, "w", encoding="utf-8").write(doc)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.goto("file://" + os.path.abspath(hp))
        pg.pdf(path=os.path.join(a.out, stem + ".pdf"), format="A4", print_background=True,
               margin=dict(top="16mm", bottom="16mm", left="15mm", right="15mm"))
        b.close()
    with open(os.path.join(a.out, stem + "_QA.txt"), "w", encoding="utf-8") as f:
        f.write(f"QA for {stem}\nStatus: {'DRAFT' if draft else 'FINAL'}\n\nBLOCKING ({len(qa)}):\n")
        f.write("\n".join(f"- {x}" for x in qa) or "- none")
        f.write("\n\nFOR JIM ONLY:\n" + "\n".join(f"- {x}" for x in info) + "\n")
    json.dump({k: dict(r, score=r["score"]) for k, r in res.items()} | {"_OBI": obi, "_SII": sii, "_T3avg": t3avg, "_completeness": completeness},
              open(os.path.join(a.out, stem + "_scores.json"), "w"), indent=1)
    print("wrote", stem, "draft" if draft else "final", "| blocking:", len(qa))


if __name__ == "__main__":
    main()
